# -*- coding: utf-8 -*-
"""Genera el documento Word: Configuración de Resend (DNS) para SpeedSkateTrack Hub."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ACCENT = RGBColor(0x0F, 0x6B, 0x5C)   # verde-teal oscuro
INK = RGBColor(0x1A, 0x1D, 0x22)
MUTED = RGBColor(0x55, 0x5B, 0x66)

doc = Document()

# Márgenes y fuente base
section = doc.sections[0]
section.left_margin = Cm(2.2)
section.right_margin = Cm(2.2)
section.top_margin = Cm(2)
section.bottom_margin = Cm(2)

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style.font.color.rgb = INK
style.paragraph_format.space_after = Pt(6)


def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def add_title(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(22)
    r.font.bold = True
    r.font.color.rgb = ACCENT
    p.paragraph_format.space_after = Pt(2)
    return p


def add_subtitle(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(12)
    r.font.color.rgb = MUTED
    r.italic = True
    p.paragraph_format.space_after = Pt(18)
    return p


def add_h1(text, num=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    if num:
        rn = p.add_run(f"{num}  ")
        rn.font.size = Pt(14)
        rn.font.bold = True
        rn.font.color.rgb = MUTED
    r = p.add_run(text)
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = ACCENT
    # línea inferior simple usando borde de párrafo
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "0F6B5C")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def add_body(text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead:
        rb = p.add_run(bold_lead + " ")
        rb.bold = True
    p.add_run(text)
    return p


def add_bullet(text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_lead:
        rb = p.add_run(bold_lead + " ")
        rb.bold = True
    p.add_run(text)
    return p


def add_callout(label, text, color_hex="E6F4F0", border_hex="0F6B5C"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    set_cell_shading(cell, color_hex)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:color"), border_hex)
        borders.append(el)
    tcPr.append(borders)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    rl = p.add_run(label + "  ")
    rl.bold = True
    rl.font.color.rgb = RGBColor(0x0F, 0x6B, 0x5C)
    p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


# ── Portada ──────────────────────────────────────────────────────────────
add_title("Configuración de Resend (DNS)")
add_subtitle("SpeedSkateTrack Hub — Guía para el administrador del dominio arkanatech.tech")

meta = doc.add_paragraph()
meta.add_run("Preparado para: ").bold = True
meta.add_run("responsable de DNS / dominios (Hostinger)\n")
meta.add_run("Dominio: ").bold = True
meta.add_run("arkanatech.tech\n")
meta.add_run("Fecha: ").bold = True
meta.add_run("agosto 2026")
meta.paragraph_format.space_after = Pt(4)

doc.add_paragraph()

# ── 1. Qué es Resend ─────────────────────────────────────────────────────
add_h1("¿Qué es Resend?", "1.")
add_body(
    "Resend es un servicio de envío de correos electrónicos transaccionales — es decir, "
    "correos automáticos que dispara una aplicación (no boletines de marketing masivo). "
    "Es la infraestructura que usa el backend de SpeedSkateTrack Hub para enviar emails "
    "de forma confiable, similar en función a servicios como SendGrid o Mailgun, pero con "
    "una configuración más simple y un panel más moderno."
)
add_body(
    "Por debajo, Resend utiliza Amazon SES (Simple Email Service), lo que le da alta "
    "entregabilidad — es decir, buena probabilidad de que los correos lleguen a la bandeja "
    "de entrada y no a spam, siempre que el dominio esté correctamente autenticado con los "
    "registros DNS que se detallan en este documento."
)

# ── 2. Cómo se usa en el proyecto ────────────────────────────────────────
add_h1("¿Cómo se usa en SpeedSkateTrack Hub?", "2.")
add_body(
    "El backend (FastAPI) llama a la API de Resend cada vez que necesita enviar un correo "
    "automático a un usuario del club. Algunos ejemplos concretos:"
)
add_bullet("Correo de bienvenida cuando se registra un nuevo atleta o usuario.")
add_bullet("Facturación mensual: envío de facturas y recordatorios de pago.")
add_bullet("Notificaciones de check-in / asistencia (ej. ingreso con NFC en la pista).")
add_bullet("Alertas y reportes automáticos dirigidos a entrenadores y administradores.")
add_bullet("Recuperación de acceso y comunicaciones de seguridad de la cuenta.")

# ── 3. Ventajas ───────────────────────────────────────────────────────────
add_h1("Ventajas de usar Resend", "3.")
add_bullet(
    "Sin servidores propios de correo (SMTP): todo se gestiona vía API, sin mantenimiento.",
    bold_lead="Simplicidad —",
)
add_bullet(
    "El plan gratuito incluye 3.000 correos/mes y 100/día — más que suficiente para el "
    "volumen esperado del club.",
    bold_lead="Capa gratuita generosa —",
)
add_bullet(
    "Al autenticar el dominio con DKIM y SPF (este documento), los correos tienen mucha "
    "menor probabilidad de caer en la carpeta de spam.",
    bold_lead="Buena entregabilidad —",
)
add_bullet(
    "Panel con estadísticas de envíos, aperturas, rebotes y quejas, útil para detectar "
    "problemas rápido.",
    bold_lead="Visibilidad —",
)
add_bullet(
    "Correos enviados desde una dirección propia (ej. noreply@arkanatech.tech) en vez de "
    "un dominio genérico — se ve más profesional y confiable para los usuarios del club.",
    bold_lead="Dominio propio —",
)

# ── 4. Qué se necesita ────────────────────────────────────────────────────
add_h1("Qué se necesita del administrador del dominio", "4.")
add_body(
    "El dominio arkanatech.tech está gestionado en Hostinger. Para que Resend pueda enviar "
    "correos en nombre de este dominio de forma autenticada, es necesario agregar 3 "
    "registros DNS en la Zona DNS de Hostinger. Estos registros NO afectan ni reemplazan "
    "los registros existentes del dominio (como los que apuntan track., stride. y split. "
    "al servidor de la aplicación) — son adicionales."
)

# Tabla de registros DNS
add_h1("Registros DNS a agregar", "5.")
rows = [
    ("Tipo", "Nombre (Name/Host)", "Contenido / Apunta a", "TTL", "Prioridad"),
    (
        "TXT",
        "resend._domainkey",
        "p=MIGfMA0GCSqGSIb3DQEBAQUAA4GNADCBiQKBgQCW4zcAFVMugCOlALmnSqpqFJkbP0GgiipdvwwPJjrZJU1COd97VNycIAdlEkRX+d4rfdGH2sNwrEWyzg/3zXjZzXuy/bsUntmrd6vZGSR4AkOx7Ph6yEibqMhEj3XaGjEUxfK76TsLKbP8UB03IliH76PYxaoQOCfNN88IvGwzswIDAQAB",
        "14400 (o el mínimo disponible)",
        "—",
    ),
    ("MX", "send", "feedback-smtp.sa-east-1.amazonses.com", "3600", "10"),
    ("TXT", "send", "v=spf1 include:amazonses.com ~all", "3600", "—"),
]
tbl = doc.add_table(rows=len(rows), cols=5)
tbl.style = "Light Grid Accent 1"
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
widths = [Cm(1.6), Cm(3.0), Cm(8.2), Cm(2.6), Cm(1.8)]
for i, row_data in enumerate(rows):
    for j, val in enumerate(row_data):
        cell = tbl.rows[i].cells[j]
        cell.width = widths[j]
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(val)
        run.font.size = Pt(9 if j == 2 and i > 0 else 10)
        if i == 0:
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            set_cell_shading(cell, "0F6B5C")
        if j == 2 and i > 0:
            run.font.name = "Consolas"

doc.add_paragraph().paragraph_format.space_after = Pt(6)

add_callout(
    "Importante:",
    "en el campo “Nombre” escribir solo la parte corta (resend._domainkey o send), NO el "
    "dominio completo. Hostinger agrega automáticamente “.arkanatech.tech” al final. Si se "
    "escribe el dominio completo, quedaría duplicado y el registro no funcionaría.",
)

# ── 6. Pasos en Hostinger ─────────────────────────────────────────────────
add_h1("Pasos en Hostinger (hPanel)", "6.")
steps = [
    "Iniciar sesión en hPanel de Hostinger.",
    "Ir a Dominios → seleccionar arkanatech.tech.",
    "Entrar a la sección “DNS / Nameservers” → “Zona DNS” (DNS Zone Editor).",
    "Por cada fila de la tabla anterior, hacer clic en “Añadir registro”.",
    "Elegir el Tipo (TXT o MX según corresponda), completar Nombre, Contenido/Apunta a, "
    "TTL y Prioridad (solo aplica al registro MX) exactamente como se detalla arriba.",
    "Guardar cada registro.",
    "No modificar ni eliminar los registros A existentes de track, stride o split — son "
    "independientes y pertenecen a la aplicación web.",
]
for i, s in enumerate(steps, 1):
    p = doc.add_paragraph(style="List Number")
    p.add_run(s)

# ── 7. Verificación ────────────────────────────────────────────────────────
add_h1("Cómo verificar que quedó correcto", "7.")
add_body(
    "La propagación de DNS puede tardar entre 15 minutos y 2 horas. Una vez transcurrido "
    "ese tiempo:"
)
add_bullet("En el panel de Resend, la sección del dominio debe mostrar los 3 registros con un check verde (✅).")
add_bullet(
    "También se puede verificar desde una terminal con estos comandos — deben devolver "
    "los mismos valores de la tabla:"
)
p = doc.add_paragraph()
p.paragraph_format.left_indent = Cm(0.6)
run = p.add_run(
    "dig TXT resend._domainkey.arkanatech.tech +short\n"
    "dig MX send.arkanatech.tech +short\n"
    "dig TXT send.arkanatech.tech +short"
)
run.font.name = "Consolas"
run.font.size = Pt(9.5)
run.font.color.rgb = MUTED

add_callout(
    "Nota final:",
    "estos registros son exclusivos para el envío de correo vía Resend y no interfieren "
    "con ningún otro servicio del dominio. Ante cualquier duda, contactar al equipo de "
    "desarrollo de SpeedSkateTrack Hub.",
    color_hex="F4E9D8",
    border_hex="A6771A",
)

out_path = r"C:\Users\MIGUEL~1\AppData\Local\Temp\claude\C--Users-Miguel-Angel-GG-Documents-PROYECTOS-CLAUDE-SKATE-skate-track-global-hub-main-skate-track-global-hub-main\0c815f62-c3f7-4974-b107-a73261cf6d11\scratchpad\Configuracion_Resend_DNS_SpeedSkateTrack.docx"
doc.save(out_path)
print("OK:", out_path)
